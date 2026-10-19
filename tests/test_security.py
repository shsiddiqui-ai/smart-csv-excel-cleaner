
"""Security and Flask integration tests."""

import io
import unittest

from openpyxl import load_workbook

from app import app


SAMPLE_CSV = (
    "Name,Email,Phone,City\n"
    "Ali Khan,ali@example.com,03001234567,Karachi\n"
)

UNSAFE_CSV = (
    "Name,Email,Phone,City\n"
    "=2+2,ali@example.com,03001234567,Karachi\n"
)


class TestSecurityAndWeb(unittest.TestCase):

    def setUp(self):
        app.config["TESTING"] = True
        self.client = app.test_client()

    def upload(self, content, filename):
        return self.client.post(
            "/clean",
            data={
                "file": (
                    io.BytesIO(content.encode("utf-8")),
                    filename
                )
            },
            content_type="multipart/form-data"
        )

    def test_home_page(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            b"Smart CSV",
            response.data
        )

    def test_valid_csv_upload(self):
        response = self.upload(
            SAMPLE_CSV,
            "contacts.csv"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

        workbook = load_workbook(
            io.BytesIO(response.data)
        )

        self.assertIn(
            "Cleaned Data",
            workbook.sheetnames
        )
        self.assertIn(
            "Quality Summary",
            workbook.sheetnames
        )

    def test_invalid_extension_rejected(self):
        response = self.upload(
            SAMPLE_CSV,
            "contacts.txt"
        )

        self.assertEqual(response.status_code, 400)

    def test_missing_columns_rejected(self):
        response = self.upload(
            "Name,Email\nAli,ali@example.com\n",
            "incomplete.csv"
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn(
            b"Missing required columns",
            response.data
        )

    def test_excel_formula_saved_as_text(self):
        response = self.upload(
            UNSAFE_CSV,
            "unsafe.csv"
        )

        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(
            io.BytesIO(response.data)
        )

        cell = workbook["Cleaned Data"]["A2"]

        # Must be an Excel text cell, not a formula cell.
        self.assertEqual(cell.data_type, "s")
        self.assertNotEqual(cell.data_type, "f")
        self.assertEqual(cell.value, "'=2+2")

    def test_oversized_upload_rejected(self):
        oversized = b"x" * (11 * 1024 * 1024)

        response = self.client.post(
            "/clean",
            data={
                "file": (
                    io.BytesIO(oversized),
                    "large.csv"
                )
            },
            content_type="multipart/form-data"
        )

        self.assertEqual(response.status_code, 413)


if __name__ == "__main__":
    unittest.main()
