
"""Smart CSV and Excel Data Cleaner - portfolio project."""

from pathlib import Path
import argparse
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REQUIRED_COLUMNS = ["Name", "Email", "Phone", "City"]


def check_email(value):
    """Check email syntax, not whether an inbox exists."""
    if not value:
        return "Missing"

    if re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
        return "Format OK (unverified)"

    return "Invalid format"


def check_phone(value):
    """Check basic phone syntax, not whether a number is active."""
    if not value:
        return "Missing"

    if re.fullmatch(r"\+?\d{7,15}", value):
        return "Format OK (unverified)"

    return "Invalid format"


def protect_spreadsheet_value(value):
    """
    Prevent spreadsheet applications from interpreting
    uploaded text as a formula.

    Protection is applied to exported values only.
    The original cleaning data stays unchanged.
    """
    if not isinstance(value, str) or not value:
        return value

    # Also check after leading whitespace or control characters.
    suspicious = value.lstrip(" \t\r\n")

    if suspicious.startswith(("=", "+", "-", "@")):
        return "'" + value

    if value.startswith(("\t", "\r", "\n")):
        return "'" + value

    return value


def prepare_safe_export(frame):
    """Create a spreadsheet-safe copy for CSV/XLSX export."""
    safe = frame.copy()

    for column in safe.columns:
        safe[column] = safe[column].map(
            protect_spreadsheet_value
        )

    return safe


def style_sheet(sheet):
    """Apply professional formatting to an Excel worksheet."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D"
    )

    header_font = Font(
        name="Aptos",
        bold=True,
        color="FFFFFF"
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    sheet.row_dimensions[1].height = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)

        lengths = [
            len(str(cell.value))
            for cell in column
            if cell.value is not None
        ]

        width = min(max(lengths, default=10) + 3, 45)

        sheet.column_dimensions[letter].width = max(
            width, 14
        )


def clean_file(source: Path, destination: Path) -> dict:
    """Clean CSV/Excel contacts and generate quality reports."""

    source = Path(source)
    destination = Path(destination)

    # 1. Validate input and output types.
    if source.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError(
            "Input must be a .csv or .xlsx file"
        )

    if destination.suffix.lower() not in (".csv", ".xlsx"):
        raise ValueError(
            "Output must be a .csv or .xlsx file"
        )

    # 2. Load input data as text.
    if source.suffix.lower() == ".csv":
        frame = pd.read_csv(
            source,
            dtype=str,
            keep_default_na=False
        )
    else:
        frame = pd.read_excel(
            source,
            dtype=str,
            keep_default_na=False
        )

    # 3. Check required columns.
    missing_columns = [
        column
        for column in REQUIRED_COLUMNS
        if column not in frame.columns
    ]

    if missing_columns:
        raise ValueError(
            "Missing required columns: "
            + ", ".join(missing_columns)
        )

    frame = frame[REQUIRED_COLUMNS].copy()
    input_rows = len(frame)

    # 4. Normalize text and contact details.
    for column in REQUIRED_COLUMNS:
        frame[column] = (
            frame[column]
            .fillna("")
            .astype(str)
            .str.strip()
        )

    frame["Name"] = frame["Name"].str.replace(
        r"\s+", " ", regex=True
    )

    frame["City"] = frame["City"].str.replace(
        r"\s+", " ", regex=True
    )

    frame["Email"] = frame["Email"].str.lower()

    frame["Phone"] = frame["Phone"].apply(
        lambda value: re.sub(r"[\s()\-]", "", value)
    )

    # 5. Validate formats.
    frame["Email_Status"] = frame["Email"].apply(
        check_email
    )

    frame["Phone_Status"] = frame["Phone"].apply(
        check_phone
    )

    # 6. Remove exact duplicates after normalization.
    duplicate_mask = frame.duplicated(
        subset=REQUIRED_COLUMNS,
        keep="first"
    )

    duplicates_removed = int(
        duplicate_mask.sum()
    )

    cleaned = frame.loc[
        ~duplicate_mask
    ].copy()

    # 7. Generate quality summary.
    quality_report = {
        "Total Input Records": input_rows,
        "Duplicate Records": duplicates_removed,
        "Unique Records": len(cleaned),
        "Missing Emails": int(
            (frame["Email_Status"] == "Missing").sum()
        ),
        "Invalid Email Formats": int(
            (frame["Email_Status"] == "Invalid format").sum()
        ),
        "Missing Phones": int(
            (frame["Phone_Status"] == "Missing").sum()
        ),
        "Invalid Phone Formats": int(
            (frame["Phone_Status"] == "Invalid format").sum()
        ),
    }

    summary = pd.DataFrame(
        quality_report.items(),
        columns=["Metric", "Count"]
    )

    # 8. Make the exported values safe for spreadsheets.
    safe_cleaned = prepare_safe_export(cleaned)

    destination.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # 9. Save cleaned output.
    if destination.suffix.lower() == ".xlsx":

        with pd.ExcelWriter(
            destination,
            engine="openpyxl"
        ) as writer:

            safe_cleaned.to_excel(
                writer,
                sheet_name="Cleaned Data",
                index=False
            )

            summary.to_excel(
                writer,
                sheet_name="Quality Summary",
                index=False
            )

            for sheet in writer.book.worksheets:
                style_sheet(sheet)

            data_sheet = writer.book["Cleaned Data"]

            warning_fill = PatternFill(
                fill_type="solid",
                fgColor="FFF2CC"
            )

            # Highlight missing/invalid status cells.
            for row in data_sheet.iter_rows(
                min_row=2,
                min_col=5,
                max_col=6
            ):
                for cell in row:
                    if cell.value != "Format OK (unverified)":
                        cell.fill = warning_fill

    else:
        safe_cleaned.to_csv(
            destination,
            index=False,
            encoding="utf-8-sig"
        )

    # 10. Save a separate quality report.
    report_path = destination.with_name(
        destination.stem + "_quality_report.csv"
    )

    summary.to_csv(
        report_path,
        index=False,
        encoding="utf-8-sig"
    )

    return {
        "input_rows": input_rows,
        "duplicates_removed": duplicates_removed,
        "output_rows": len(cleaned),
        "output_path": str(destination),
        "quality_report_path": str(report_path),
    }


if __name__ == "__main__":

    parser = argparse.ArgumentParser(
        description=(
            "Clean CSV/Excel contacts "
            "and generate quality reports"
        )
    )

    parser.add_argument(
        "--input",
        type=Path,
        default=Path("input/sample_contacts.csv")
    )

    parser.add_argument(
        "--output",
        type=Path,
        default=Path("output/cleaned_contacts.xlsx")
    )

    args = parser.parse_args()

    try:
        result = clean_file(
            args.input,
            args.output
        )

    except (
        ValueError,
        FileNotFoundError,
        ImportError,
        PermissionError
    ) as exc:
        parser.exit(
            1,
            f"Error: {exc}\n"
        )

    print("Cleaning complete!")

    for key, value in result.items():
        print(f"{key}: {value}")
